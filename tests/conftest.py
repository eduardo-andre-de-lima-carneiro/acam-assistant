from __future__ import annotations

import textwrap
from pathlib import Path

import docx
import pytest

from acam import storage

from pdfgen import make_pdf

ROOT = Path(__file__).resolve().parents[1]
SAMPLE_MD = ROOT / "examples" / "sample-article.md"


@pytest.fixture(autouse=True)
def _isolated_env(monkeypatch):
    """Tests control engines and workers explicitly; ignore the developer's environment."""
    for var in ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN", "ANTHROPIC_PROFILE",
                "ACAM_ALLOW_SIGNUP", "ACAM_SECURE_COOKIES", "ACAM_ADMIN_USER",
                "ACAM_ADMIN_PASSWORD", "ACAM_ADMIN_PASSWORD_FILE"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("ACAM_EMBEDDED_WORKER", "false")


@pytest.fixture
def workspace(tmp_path: Path) -> Path:
    return tmp_path / "projects"


@pytest.fixture
def project(workspace: Path) -> storage.Project:
    return storage.create_project(workspace, "AI in Software Engineering",
                                  research_questions=["How is AI applied to code review?"])


@pytest.fixture
def sample_md(tmp_path: Path) -> Path:
    dst = tmp_path / "sample-article.md"
    dst.write_text(SAMPLE_MD.read_text(encoding="utf-8"), encoding="utf-8")
    return dst


def _plain_lines() -> list[str]:
    """Sample article as plain text lines (no Markdown syntax)."""
    out = []
    for ln in SAMPLE_MD.read_text(encoding="utf-8").split("\n"):
        ln = ln.lstrip("#> ").strip()
        out.extend(textwrap.wrap(ln, 95) or [""])
    return out


@pytest.fixture
def sample_txt(tmp_path: Path) -> Path:
    p = tmp_path / "sample-article.txt"
    p.write_text("\n".join(_plain_lines()), encoding="utf-8")
    return p


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    lines = _plain_lines()
    pages = [lines[i:i + 50] for i in range(0, len(lines), 50)]
    p = tmp_path / "sample-article.pdf"
    p.write_bytes(make_pdf(pages, title="Large Language Models as Code Review Assistants",
                           author="Maria Fernanda Souza, John Peter Miller"))
    return p


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    d = docx.Document()
    d.core_properties.title = "Large Language Models as Code Review Assistants"
    d.core_properties.author = "Maria Fernanda Souza; John Peter Miller"
    for ln in SAMPLE_MD.read_text(encoding="utf-8").split("\n"):
        if not ln.strip() or ln.startswith(">"):
            continue
        if ln.startswith("# "):
            d.add_paragraph(ln[2:], style="Title")
        elif ln.startswith("## "):
            d.add_heading(ln[3:], level=1)
        else:
            d.add_paragraph(ln)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Metric", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "Review time reduction", "23%"
    p = tmp_path / "sample-article.docx"
    d.save(str(p))
    return p


@pytest.fixture
def sample_html(tmp_path: Path) -> Path:
    p = tmp_path / "page.html"
    p.write_text("""<!doctype html><html><head><title>Cidades</title>
<meta name="author" content="Instituto Brasileiro de Geografia e Estatística">
<script>var x = 1;</script><style>p{}</style></head>
<body><nav>Menu Home</nav><h1>Cidades e Estados</h1>
<p>Portal with <strong>statistical</strong> information about Brazilian municipalities.
See <a href="https://cidades.ibge.gov.br">the portal</a>.</p>
<h2>Methodology</h2><p>Data are compiled from the national census and administrative records,
updated annually by the institute across all municipalities in the country.</p>
<table><tr><th>City</th><th>Population</th></tr><tr><td>Recife</td><td>1,488,920</td></tr></table>
<h2>Conclusion</h2><p>The portal consolidates official statistics for public use and research.</p>
</body></html>""", encoding="utf-8")
    return p
